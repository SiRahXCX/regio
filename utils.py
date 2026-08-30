import csv
from pypdf import PdfReader, PdfWriter
from openpyxl import load_workbook


def fill_pdf_template_from_csv(csv_filepath: str, template_filepath: str, output_filepath: str) -> None:
    with open(csv_filepath, mode='r', encoding='utf-8-sig') as csv_file:
        reader = csv.DictReader(csv_file)

        for idx, row in enumerate(reader):
            pdf_reader = PdfReader(template_filepath)
            pdf_writer = PdfWriter()
            pdf_writer.append(pdf_reader)

            pdf_writer.update_page_form_field_values(
                pdf_writer.pages[0],
                row
            )

            output_filename = f'{output_filepath}/output_{idx+1}.pdf'
            with open(output_filename, mode='wb') as output_file:
                pdf_writer.write(output_file)


def fill_pdf_template_from_xlsx(xlsx_filepath: str, template_filepath: str, output_filepath: str) -> None:
    workbook = load_workbook(xlsx_filepath, read_only=True, data_only=True)
    worksheet = workbook.active 
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    reader = [dict(zip(headers, row)) for row in rows] 

    for idx, row in enumerate(reader):
        pdf_reader = PdfReader(template_filepath)
        pdf_writer = PdfWriter()
        pdf_writer.append(pdf_reader)

        pdf_writer.update_page_form_field_values(
            pdf_writer.pages[0],
            row
        )

        output_filename = f'{output_filepath}/output_{idx+1}.pdf'
        with open(output_filename, mode='wb') as output_file:
            pdf_writer.write(output_file)
